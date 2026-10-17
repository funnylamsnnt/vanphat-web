#!/usr/bin/env python3
"""Export Danh Mục Vật Tư from Vạn Phát Excel → data/products.json.

Maps each Excel nhóm onto lĩnh vực using data/ia.json (config array, not a
fixed count of domains). Placeholder nhóm such as Bánh stay in the catalog
even when Excel has zero SKUs.

Images already stored for a mã are kept when the Excel cell is empty.
Customer order-form / Apps Script links are not written.

After a successful export this script regenerates indexable category, lĩnh vực,
and product HTML plus sitemap.xml and js/catalog-paths.js.

    python3 scripts/sync_products_from_excel.py path/to.xlsx \
        --drive-title "<title from Drive get_file_metadata of 1xgIKFgdd1iki3uoeIRh29FMrOzxRfreR>"

SOURCE TITLE GUARD (10/10/2026): --drive-title is required. If the Drive title
ends with .tmp or is not "PHIEU BAO GIA VAN PHAT - 2026.xlsx" (normalised), the
script exits 3 before writing anything — do not deploy, do not refresh the LIVE
bridge, notify anh Hà. See scripts/source_title_guard.py.
"""
import argparse, json, re, hashlib, subprocess, sys, unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
IA_PATH = ROOT / "data" / "ia.json"

# Default (legacy) layout: A nhóm, B mã, C tên, D ĐVT, E giá, F link ảnh, K tồn.
# Thương Hiệu and Quy Cách are not here: they are found by header name only.
DEFAULT_COLS = {"nhom": 0, "ma": 1, "ten": 2, "dvt": 3, "gia": 4, "anh": 5, "ton": 10}
# Optional columns. A missing header does not fail the export and never falls
# back to a fixed index such as column S or T.
OPTIONAL_HEADER_KEYS = ("thuong_hieu", "quy_cach")
HEADER_PATTERNS = {
    "nhom": [r"^nh[oó]m"],
    "ma": [r"^m[aã]\s*(sp|s[aả]n ph[aẩ]m|h[aà]ng|vt)?$"],
    "ten": [r"^t[eê]n\s*(s[aả]n ph[aẩ]m|sp|h[aà]ng|m[aặ]t h[aà]ng|v[aậ]t t[uư])"],
    "dvt": [r"^đvt$", r"^đ[oơ]n v[iị] t[ií]nh"],
    "gia": [r"^đ[oơ]n gi[aá]", r"^gi[aá] b[aá]n", r"^gi[aá]$"],
    "anh": [r"(link|url).*(h[iì]nh|[aả]nh)", r"^h[iì]nh [aả]nh", r"^[aả]nh$"],
    "ton": [r"^t[oồ]n kho", r"^t[oồ]n$"],
    "thuong_hieu": [r"^th[uư][oơ]ng hi[eệ]u"],
    "quy_cach": [r"^quy c[aá]ch"],
}


def detect_columns(header) -> dict:
    """Map fields to column indexes from the header row; fall back to the legacy layout.

    thuong_hieu and quy_cach are optional. They are set only when the header
    matches, even if the required columns fall back to DEFAULT_COLS.
    """
    cols = dict(DEFAULT_COLS)
    names = [re.sub(r"\s+", " ", str(h or "")).strip().lower() for h in header]
    found = {}
    for key, pats in HEADER_PATTERNS.items():
        for idx, name in enumerate(names):
            if name and any(re.search(p, name) for p in pats):
                found[key] = idx
                break
    required = {key: idx for key, idx in found.items() if key not in OPTIONAL_HEADER_KEYS}
    optional = {key: idx for key, idx in found.items() if key in OPTIONAL_HEADER_KEYS}
    if all(k in required for k in ("ma", "ten", "gia")):
        cols.update(required)
    for key in OPTIONAL_HEADER_KEYS:
        cols.pop(key, None)
    cols.update(optional)
    return cols


def clean(v) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def fold_name(value: str) -> str:
    s = clean(value).lower().replace("đ", "d")
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def slugify(text: str, max_len: int = 60) -> str:
    s = str(text or "").strip().lower().replace("đ", "d")
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    if max_len:
        s = s[:max_len].strip("-")
    return s


def to_int(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    s = str(v).strip().replace("₫", "").replace("đ", "").replace(" ", "")
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", s):
        s = re.sub(r"[.,]", "", s)
    try:
        return int(round(float(s.replace(",", "."))))
    except Exception:
        return None


def load_ia() -> dict:
    if not IA_PATH.is_file():
        raise SystemExit(f"Missing lĩnh vực config: {IA_PATH}")
    data = json.loads(IA_PATH.read_text(encoding="utf-8"))
    linh = data.get("linh_vuc") or []
    nhom = data.get("nhom") or []
    if not linh or not nhom:
        raise SystemExit("data/ia.json needs linh_vuc[] and nhom[]")
    return data


DEFAULT_EXCLUDE_GROUPS = ["Không đưa lên web"]


def nhom_key(value: str) -> str:
    """Trim, drop case/diacritics, and treat hyphens as spaces. Blank stays blank."""
    s = clean(value).replace("\u00a0", " ").replace("\u200b", "").replace("\ufeff", "")
    s = fold_name(s).replace("-", " ").replace("_", " ").replace("/", " ")
    return re.sub(r"\s+", " ", s).strip(" .:-")


def exclude_group_folds(ia: dict) -> set:
    """Keys for Excel nhóm that must never be published.

    data/ia.json excludeGroups, plus "Không đưa lên web" even if that key is
    missing. Matching ignores case, accents (including đ/Đ), extra spaces and
    hyphens, and allows a suffix such as "(nội bộ)".
    """
    names = list(DEFAULT_EXCLUDE_GROUPS) + list(ia.get("excludeGroups") or [])
    return {nhom_key(n) for n in names if nhom_key(n)}


def is_excluded_nhom(raw_nhom: str, phrases: set) -> bool:
    folded = nhom_key(raw_nhom)
    if not folded:
        return False
    padded = f" {folded} "
    for phrase in phrases:
        if folded == phrase or folded.startswith(phrase + " ") or f" {phrase} " in padded:
            return True
    return False


# Ô trống và nhãn «Khác» không có lĩnh vực khớp. Nhóm có tên thì tự xếp.
DRINK_TOKENS = frozenset({
    "tra", "sua", "banh", "snack", "soda", "cafe", "coffee",
    "sting", "pepsi", "coca", "mirinda", "aquafina", "lavie", "dasani",
})
DRINK_PHRASES = (
    "nuoc suoi", "nuoc ngot", "nuoc uong", "nuoc tang luc", "nuoc khoang",
    "nuoc tinh khiet", "do uong", "ca phe", "do an", "banh keo",
)
DIEN_PHRASES = (
    "o cam", "day dien", "thiet bi dien", "dien gia dung", "o dien",
    "phich cam", "cong tac", "bong den", "may lanh", "quat dien",
    "noi com", "noi chien",
)


def is_hold_label(raw_nhom: str) -> bool:
    """Blank nhóm and the junk label Khác stay off the web."""
    key = nhom_key(raw_nhom)
    if not key:
        return True
    return key == "khac" or key.startswith("khac ") or key.startswith("khac(")


def classify_new_nhom(raw_nhom: str) -> tuple[str | None, str]:
    """Pick one of the three lĩnh vực, or None when the nhóm must stay off the web.

    A new L2 is created under the chosen L1. Nothing is dumped into Dụng cụ VP.
    """
    if is_hold_label(raw_nhom):
        if not nhom_key(raw_nhom):
            return None, "ô Nhóm trống"
        return None, "nhãn Khác, không có lĩnh vực khớp"
    folded = fold_name(raw_nhom)
    tokens = set(re.findall(r"[a-z0-9]+", folded))
    drink = 0
    if tokens & DRINK_TOKENS:
        drink += 2
    for phrase in DRINK_PHRASES:
        if phrase in folded:
            drink += 3
    dien = 0
    if "dien" in tokens:
        dien += 3
    for phrase in DIEN_PHRASES:
        if phrase in folded:
            dien += 3
    if drink > dien and drink > 0:
        return "nuoc-banh", "tên khớp nước uống hoặc bánh"
    if dien > drink and dien > 0:
        return "dien-gia-dung", "tên khớp điện gia dụng"
    if drink > 0 and drink == dien:
        return "vpp", "khớp cả nước uống và điện, tạm xếp Văn phòng phẩm để anh chuyển nếu cần"
    return "vpp", "không khớp Nước uống hay Điện, xếp Văn phòng phẩm bằng nhóm L2 mới"


def _unique_slug(base: str, nhom_cfg: list) -> str:
    used = set()
    for item in nhom_cfg:
        used.add(item.get("id") or "")
        used.add(item.get("slug") or "")
    candidate = base or "nhom"
    if candidate not in used:
        return candidate
    number = 2
    while f"{candidate}-{number}" in used:
        number += 1
    return f"{candidate}-{number}"


def ensure_group(raw_nhom: str, ia: dict, nhom_cfg: list, by_fold: dict, lv_by_id: dict, placed: dict) -> dict:
    """Map a named nhóm onto a lĩnh vực, persist it on the in-memory ia, and publish it."""
    key = fold_name(raw_nhom)
    current = by_fold.get(key)
    if current and lv_by_id.get(current.get("linhVucId") or ""):
        return current
    lv_id, reason = classify_new_nhom(raw_nhom)
    if not lv_id:
        raise SystemExit(f"Nhóm {raw_nhom!r} phải giữ ngoài web, không gọi ensure_group")
    lv = lv_by_id.get(lv_id)
    if not lv:
        raise SystemExit(f"Thiếu lĩnh vực {lv_id} trong data/ia.json")
    if current:
        meta = current
        meta["linhVucId"] = lv_id
        meta["autoPlaced"] = True
        for original in ia.get("nhom") or []:
            if fold_name(original.get("TenNhom") or "") == key:
                original["linhVucId"] = lv_id
                original["autoPlaced"] = True
    else:
        slug = _unique_slug(slugify(raw_nhom) or "nhom", nhom_cfg)
        meta = {
            "id": slug,
            "STT": 1000 + len(placed),
            "TenNhom": raw_nhom,
            "slug": slug,
            "linhVucId": lv_id,
            "visible": True,
            "moTa": f"Nhóm {raw_nhom}, tự xếp vào {lv.get('ten') or lv_id}.",
            "autoPlaced": True,
        }
        nhom_cfg.append(meta)
        ia.setdefault("nhom", []).append(meta)
        by_fold[key] = meta
    sid = meta.get("id") or meta.get("slug")
    ids = lv.setdefault("nhomIds", [])
    if sid and sid not in ids:
        ids.append(sid)
    placed.setdefault(
        meta.get("TenNhom") or raw_nhom,
        {
            "linhVucId": lv_id,
            "linhVucTen": lv.get("ten") or "",
            "lyDo": reason,
            "ma": [],
        },
    )
    return meta


def assert_optional_columns() -> None:
    """Thương Hiệu and Quy Cách follow the header. No fixed column fallback."""
    full = [
        "Nhóm Hàng", "Mã SP", "Tên Sản Phẩm", "ĐVT", "Đơn Giá",
        "Link Hình Ảnh", "", "", "", "", "Tồn Kho",
    ]
    full += [""] * 7
    full += ["Thương Hiệu", "Quy Cách"]
    cols = detect_columns(full)
    if cols.get("thuong_hieu") != 18 or cols.get("quy_cach") != 19:
        raise SystemExit(f"header columns missed: {cols.get('thuong_hieu')} {cols.get('quy_cach')}")
    bare = ["Nhóm", "Mã SP", "Tên sản phẩm", "ĐVT", "Giá", "Ảnh", "", "", "", "", "Tồn"]
    bare += ["x"] * 10
    legacy = detect_columns(bare)
    if "thuong_hieu" in legacy or "quy_cach" in legacy:
        raise SystemExit(f"optional columns fell back without a header: {legacy}")
    shifted = ["Thuong hieu", "Quy cach", "Nhóm", "Mã SP", "Tên sản phẩm", "ĐVT", "Đơn giá"]
    shifted_cols = detect_columns(shifted)
    if shifted_cols.get("thuong_hieu") != 0 or shifted_cols.get("quy_cach") != 1:
        raise SystemExit(f"optional columns must follow the header: {shifted_cols}")
    if shifted_cols.get("ma") != 3 or shifted_cols.get("ten") != 4:
        raise SystemExit(f"required columns should still map from the header: {shifted_cols}")


def assert_placement_rules() -> None:
    expect = {
        "": None,
        "   ": None,
        "Khác": None,
        "  khác (nội bộ) ": None,
        "Trà đóng chai": "nuoc-banh",
        "Bánh quy": "nuoc-banh",
        "Ổ cắm thông minh": "dien-gia-dung",
        "Dây điện": "dien-gia-dung",
        "Bìa nhựa": "vpp",
        "Nước rửa kính": "vpp",
    }
    for raw, lv_id in expect.items():
        got, _reason = classify_new_nhom(raw)
        if got != lv_id:
            raise SystemExit(f"placement {raw!r} -> {got}, expected {lv_id}")
    if not is_hold_label("Khác") or not is_hold_label(""):
        raise SystemExit("ô trống và Khác phải được giữ")
    if is_hold_label("Bìa nhựa"):
        raise SystemExit("Bìa nhựa không phải nhãn giữ")
    ia = {
        "autoPlaceNewGroups": True,
        "linh_vuc": [
            {"id": "vpp", "ten": "Văn phòng phẩm", "nhomIds": ["dung-cu-vp"]},
            {"id": "nuoc-banh", "ten": "Nước uống & Bánh", "nhomIds": ["nuoc-uong"]},
            {"id": "dien-gia-dung", "ten": "Điện gia dụng", "nhomIds": ["dien"]},
        ],
        "nhom": [{
            "id": "dung-cu-vp",
            "TenNhom": "Dụng cụ VP",
            "slug": "dung-cu-vp",
            "linhVucId": "vpp",
        }],
    }
    nhom_cfg = [dict(item) for item in ia["nhom"]]
    by_fold = {fold_name(item["TenNhom"]): item for item in nhom_cfg}
    lv_by_id = {item["id"]: item for item in ia["linh_vuc"]}
    placed: dict = {}
    meta = ensure_group("Trà đào", ia, nhom_cfg, by_fold, lv_by_id, placed)
    if meta["linhVucId"] != "nuoc-banh" or meta["id"] == "dung-cu-vp":
        raise SystemExit(f"Trà đào bị xếp sai: {meta}")
    if meta["id"] not in lv_by_id["nuoc-banh"]["nhomIds"]:
        raise SystemExit("nhomIds của Nước uống chưa nhận nhóm mới")
    if meta["id"] in lv_by_id["vpp"]["nhomIds"]:
        raise SystemExit("không nhét nhóm nước vào Văn phòng phẩm")
    again = ensure_group("Trà đào", ia, nhom_cfg, by_fold, lv_by_id, placed)
    if again is not meta and again.get("id") != meta.get("id"):
        raise SystemExit("gọi lại ensure_group không được tạo nhóm thứ hai")


def assert_exclusion_rules() -> None:
    phrases = {nhom_key("Không đưa lên web")}
    for sample in (
        "Không đưa lên web",
        "  không đưa lên web ",
        "KHÔNG ĐƯA LÊN WEB",
        "Không đưa lên web (nội bộ)",
        "Không-đưa-lên-web",
    ):
        if not is_excluded_nhom(sample, phrases):
            raise SystemExit(f"exclusion missed {sample!r}")
    for sample in ("Giấy", "Khác", "", "  "):
        if is_excluded_nhom(sample, phrases):
            raise SystemExit(f"exclusion false positive {sample!r}")
    if nhom_key("  ") or nhom_key("Khác") == "":
        raise SystemExit("blank nhóm must not be renamed before the exclusion check")


def load_previous(out_json: Path) -> dict:
    if not out_json.exists():
        return {}
    try:
        return json.loads(out_json.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def export(xlsx_path: Path, out_json: Path) -> dict:
    from openpyxl import load_workbook

    ia = load_ia()
    prev = load_previous(out_json)
    prev_by_ma = {p.get("ma"): p for p in (prev.get("sanpham") or []) if p.get("ma")}

    nhom_cfg = [dict(n) for n in ia["nhom"]]
    by_fold = {fold_name(n.get("TenNhom")): n for n in nhom_cfg if n.get("TenNhom")}
    lv_by_id = {lv.get("id"): lv for lv in ia["linh_vuc"] if lv.get("id")}
    excl_folds = exclude_group_folds(ia)
    assert_exclusion_rules()
    assert_placement_rules()
    auto_place = ia.get("autoPlaceNewGroups", True)
    excluded = {}  # raw nhóm -> [mã] never published
    held = {}  # ô trống / Khác, hoặc nhóm tên khi autoPlaceNewGroups tắt
    placed = {}  # TenNhom -> {linhVucId, linhVucTen, lyDo, ma}

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)

    company = {
        "Ten": "CÔNG TY TNHH TƯ VẤN ĐẦU TƯ THƯƠNG MẠI VẠN PHÁT",
        "DiaChi": "LK 19-06 Đường số 20 KĐT Mỹ Gia, Vĩnh Thái, Phường Nam Nha Trang",
        "Website": "http://vanphatcompany.vn",
        "Email": "congtytnhhvanphat999@gmail.com",
        "Hotline": "033 5652 832",
        "STK": "4703201014329",
        "NganHang": "Ngân hàng Nông nghiệp và PTNT (Agribank)",
    }
    if "Đơn Đặt Hàng" in wb.sheetnames:
        ws0 = wb["Đơn Đặt Hàng"]
        for row in ws0.iter_rows(min_row=1, max_row=1, values_only=True):
            header = row[1] if row else None
            break
        if isinstance(header, str):
            for line in header.split("\n"):
                line = line.strip()
                if line.startswith("Địa chỉ:"):
                    company["DiaChi"] = line.replace("Địa chỉ:", "").strip()
                if "Email:" in line:
                    m = re.search(r"Email:\s*(\S+)", line)
                    if m:
                        company["Email"] = m.group(1).strip()
                if "STK" in line and ":" in line:
                    m = re.search(r":\s*(\d+)", line)
                    if m:
                        company["STK"] = m.group(1)

    sheet = "Danh Mục Vật Tư"
    if sheet not in wb.sheetnames:
        raise SystemExit(f"Missing sheet: {sheet}. Have: {wb.sheetnames}")

    products = []
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    col = detect_columns(header)
    for row in rows:
        if not row:
            continue
        cell = lambda k: row[col[k]] if col.get(k) is not None and col[k] < len(row) else None
        ma = clean(cell("ma"))
        if not ma:
            continue
        # Decide on the raw cell. Never rewrite a blank or an exclusion label to "Khác"
        # first — that alias is why VP415–VP449 were reported as nhóm Khác.
        raw_nhom = clean(cell("nhom"))
        if is_excluded_nhom(raw_nhom, excl_folds):
            excluded.setdefault(raw_nhom, []).append(ma)
            continue
        if is_hold_label(raw_nhom):
            held.setdefault(raw_nhom or "(trống)", []).append(ma)
            continue
        meta = by_fold.get(fold_name(raw_nhom))
        if meta is None or not lv_by_id.get(meta.get("linhVucId") or ""):
            if not auto_place:
                held.setdefault(raw_nhom, []).append(ma)
                continue
            meta = ensure_group(raw_nhom, ia, nhom_cfg, by_fold, lv_by_id, placed)
        if meta.get("TenNhom") in placed and ma not in placed[meta["TenNhom"]]["ma"]:
            placed[meta["TenNhom"]]["ma"].append(ma)
        lv_id = meta.get("linhVucId") or ""
        lv = lv_by_id.get(lv_id) or {}
        ten = clean(cell("ten"))
        dvt = clean(cell("dvt"))
        gia = to_int(cell("gia"))
        gia = gia if gia is not None and gia >= 0 else 0
        anh = clean(cell("anh"))
        if not anh:
            old = prev_by_ma.get(ma) or {}
            anh = clean(old.get("anh"))
        ton = to_int(cell("ton"))
        record = {
            "nhom": meta["TenNhom"],
            "nhomId": meta.get("id") or meta.get("slug") or "",
            "linhVucId": lv_id,
            "linhVucTen": lv.get("ten") or "",
            "ma": ma,
            "ten": ten,
            "dvt": dvt,
            "gia": gia,
            "anh": anh,
            "ton": ton if ton is not None else 0,
        }
        thuong_hieu = clean(cell("thuong_hieu"))
        quy_cach = clean(cell("quy_cach"))
        if thuong_hieu:
            record["thuongHieu"] = thuong_hieu
        if quy_cach:
            record["quyCach"] = quy_cach
        products.append(record)
    wb.close()

    by = {}
    for p in products:
        by[p["ma"]] = p
    products = list(by.values())
    leaked = [p["ma"] for p in products if is_excluded_nhom(p.get("nhom") or "", excl_folds)]
    if leaked:
        raise SystemExit(f"Excluded nhóm leaked into web catalog: {leaked}")

    stt_of = {n.get("TenNhom"): n.get("STT") or 999 for n in nhom_cfg}

    def key(p):
        m = re.search(r"(\d+)", p["ma"])
        return (stt_of.get(p["nhom"], 999), int(m.group(1)) if m else 0, p["ma"])

    products.sort(key=key)
    nhom_cfg.sort(key=lambda n: (n.get("STT") or 999, n.get("TenNhom") or ""))
    if placed:
        ia["nhom"] = sorted(ia.get("nhom") or [], key=lambda n: (n.get("STT") or 999, n.get("TenNhom") or ""))
        IA_PATH.write_text(json.dumps(ia, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    linh_vuc = sorted(ia["linh_vuc"], key=lambda lv: (lv.get("sort") or 0, lv.get("ten") or ""))
    added = [p["ma"] for p in products if p["ma"] not in prev_by_ma]
    removed = [ma for ma in prev_by_ma if ma not in by]
    price_changes = []
    slug_changes = []
    zero_price = []
    for product in products:
        if not product.get("gia"):
            zero_price.append(product["ma"])
        old = prev_by_ma.get(product["ma"])
        if not old:
            continue
        if old.get("gia") != product.get("gia"):
            price_changes.append({"ma": product["ma"], "old": old.get("gia"), "new": product.get("gia")})
        if slugify(str(old.get("ten") or "")) != slugify(str(product.get("ten") or "")):
            slug_changes.append({"ma": product["ma"], "old": old.get("ten"), "new": product.get("ten")})

    # Never emit a customer order-form URL. Ordering stays cart + hotline/Zalo.
    out = {
        "company": company,
        "linh_vuc": linh_vuc,
        "nhom": nhom_cfg,
        "sanpham": products,
        "hotline": "033 5652 832",
        "syncedFrom": xlsx_path.name,
    }
    if "orderUrl" in out:
        raise SystemExit("orderUrl must not be written")

    out_json.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(out, ensure_ascii=False, indent=2)
    prev_text = out_json.read_text(encoding="utf-8") if out_json.exists() else ""
    changed = hashlib.sha256(text.encode()).hexdigest() != hashlib.sha256(prev_text.encode()).hexdigest()
    if changed:
        out_json.write_text(text, encoding="utf-8")

    with_img = sum(1 for p in products if p["anh"] and "placeholder" not in p["anh"].lower())
    counts = {}
    for p in products:
        counts[p["nhom"]] = counts.get(p["nhom"], 0) + 1
    lv_counts = {}
    for p in products:
        lv_counts[p.get("linhVucId") or ""] = lv_counts.get(p.get("linhVucId") or "", 0) + 1

    waters_in_tools = [
        p["ma"]
        for p in products
        if p["nhom"] == "Dụng cụ VP" and fold_name(p["ten"]).find("nuoc suoi") >= 0
    ]
    if waters_in_tools:
        raise SystemExit(f"Nước uống still classified as Dụng cụ VP: {waters_in_tools}")
    missing_new = [f"VP{n}" for n in range(391, 408) if f"VP{n}" not in by]
    if missing_new:
        print("warning: expected SKUs missing from Excel:", ", ".join(missing_new))

    return {
        "changed": changed,
        "count": len(products),
        "with_img": with_img,
        "groups": counts,
        "linhVuc": lv_counts,
        "added": added,
        "removed": removed,
        "priceChanges": price_changes,
        "slugChanges": slug_changes,
        "zeroPrice": zero_price,
        "unknownGroups": sorted(placed),
        "autoPlacedGroups": [
            {"ten": name, **info} for name, info in sorted(placed.items())
        ],
        "excludedGroups": excluded,
        "heldUnmappedGroups": held,
        "keptImages": sum(1 for p in products if p["anh"] and not (prev_by_ma.get(p["ma"]) or {}).get("anh") and p["ma"] not in prev_by_ma) ,
        "withThuongHieu": sum(1 for p in products if p.get("thuongHieu")),
        "withQuyCach": sum(1 for p in products if p.get("quyCach")),
        "path": str(out_json),
    }


def regenerate_catalog_pages(products_json: Path) -> None:
    """Rebuild static SEO pages so new SKUs are crawlable after catalog sync."""
    gen = Path(__file__).resolve().parent / "generate_catalog_pages.py"
    if not gen.is_file():
        raise SystemExit(f"Missing generator: {gen}")
    subprocess.check_call([sys.executable, str(gen), "--products", str(products_json)])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("-o", "--out", default=str(ROOT / "data" / "products.json"))
    ap.add_argument("--drive-title", required=True,
                    help="Current Drive title of source fileId 1xgIKFgdd1iki3uoeIRh29FMrOzxRfreR")
    args = ap.parse_args()
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from source_title_guard import check_title, ABORT_EXIT
    ok, why = check_title(args.drive_title)
    if not ok:
        print(json.dumps({"aborted": True, "reason": why, "notify": "anh Hà"}, ensure_ascii=False))
        sys.exit(ABORT_EXIT)
    out = Path(args.out)
    stats = export(Path(args.xlsx), out)
    print(json.dumps(stats, ensure_ascii=False))
    regenerate_catalog_pages(out)


assert_optional_columns()
assert_placement_rules()

if __name__ == "__main__":
    main()
