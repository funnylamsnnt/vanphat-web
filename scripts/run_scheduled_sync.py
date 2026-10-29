#!/usr/bin/env python3
"""Lịch 14:35: tải Excel, sync catalog, đẩy main, đợi Pages, làm mới cầu nối LIVE.

Secret đọc từ môi trường. Không in và không ghi fileId cầu nối, không sửa file Excel gốc.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import quote

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from sync_schedule import (  # noqa: E402
    CHINH_SACH_TITLE,
    DANH_MUC_TITLE,
    SOURCE_FILE_ID,
    TEMP_TITLE,
    Job,
    bridge_sheet_plan,
    cell_for_sheet,
    col_letter,
    ma_from_column,
    ma_list,
    parse_sa,
    path_allowed,
    redact,
    run_job,
    source_download_request,
    source_metadata_request,
)

SCOPES = (
    "https://www.googleapis.com/auth/drive.readonly",
    "https://www.googleapis.com/auth/spreadsheets",
)
GIT_ADD = [
    "data",
    "linh-vuc",
    "san-pham",
    "dia-phuong",
    "app",
    "nhan-vien",
    "nhan-vien.html",
    "js/catalog-paths.js",
    "sitemap.xml",
    "robots.txt",
    "scripts/sync_state.json",
    "index.html",
    "san-pham.html",
    "thu-ngo.html",
    "photocopy.html",
    "lien-he.html",
    "khu-vuc.html",
    "offline.html",
]


def _session(info: dict):
    from google.auth.transport.requests import AuthorizedSession
    from google.oauth2 import service_account

    creds = service_account.Credentials.from_service_account_info(info, scopes=SCOPES)
    return AuthorizedSession(creds)


def _brief(resp, secret: str, label: str) -> str:
    body = redact(getattr(resp, "text", "") or "", [secret])[:300]
    return f"{label}, HTTP {getattr(resp, 'status_code', '?')} {body}"


def fetch_metadata(session) -> dict:
    method, url, params = source_metadata_request(SOURCE_FILE_ID)
    if method != "GET":
        raise SystemExit("chỉ được đọc nguồn bằng GET")
    resp = session.get(url, params=params, timeout=60)
    if resp.status_code != 200:
        raise SystemExit(_brief(resp, "", "Không đọc metadata Drive"))
    return resp.json()


def download_source(session, meta: dict, dest: Path) -> None:
    method, url, params = source_download_request(SOURCE_FILE_ID, str(meta.get("mimeType") or ""))
    if method != "GET":
        raise SystemExit("chỉ được tải nguồn bằng GET")
    resp = session.get(url, params=params, timeout=180)
    if resp.status_code != 200:
        raise SystemExit(_brief(resp, "", "Không tải Excel"))
    if not resp.content:
        raise SystemExit("File Excel tải về rỗng")
    dest.write_bytes(resp.content)


def read_matrix(path: Path) -> list[list]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=False, read_only=True)
    try:
        if DANH_MUC_TITLE not in wb.sheetnames:
            raise SystemExit("Excel thiếu sheet Danh Mục Vật Tư")
        ws = wb[DANH_MUC_TITLE]
        rows = [[cell_for_sheet(cell) for cell in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    while rows and all(cell == "" for cell in rows[-1]):
        rows.pop()
    return rows


def _sheets_get(session, spreadsheet_id: str) -> list[dict]:
    resp = session.get(
        f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}",
        params={"fields": "sheets(properties(sheetId,title))"},
        timeout=60,
    )
    if resp.status_code != 200:
        raise SystemExit(_brief(resp, spreadsheet_id, "Không đọc cầu nối"))
    found = []
    for sheet in resp.json().get("sheets") or []:
        prop = sheet.get("properties") or {}
        found.append({"sheetId": prop.get("sheetId"), "title": prop.get("title") or ""})
    return found


def _sheets_batch(session, spreadsheet_id: str, requests: list[dict]) -> dict:
    if not requests:
        return {}
    resp = session.post(
        f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}:batchUpdate",
        json={"requests": requests},
        timeout=120,
    )
    if resp.status_code != 200:
        raise SystemExit(_brief(resp, spreadsheet_id, "Không sửa cầu nối"))
    return resp.json()


def _delete_requests(ids: list[int], keep: list[int]) -> list[dict]:
    if set(ids) & set(keep):
        raise SystemExit("Không xóa tab Chính Sách Đối Tác")
    return [{"deleteSheet": {"sheetId": sheet_id}} for sheet_id in ids]


def _write_values(session, spreadsheet_id: str, title: str, rows: list[list]) -> None:
    range_name = f"'{title}'!A1"
    url = (
        f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}/values/"
        + quote(range_name, safe="")
    )
    resp = session.put(
        url,
        params={"valueInputOption": "USER_ENTERED"},
        json={"range": range_name, "majorDimension": "ROWS", "values": rows},
        timeout=180,
    )
    if resp.status_code != 200:
        raise SystemExit(_brief(resp, spreadsheet_id, "Không ghi tab Danh Mục"))


def _read_ma_column(session, spreadsheet_id: str, title: str, index: int, n_rows: int) -> list:
    letter = col_letter(index)
    range_name = f"'{title}'!{letter}1:{letter}{n_rows}"
    url = (
        f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}/values/"
        + quote(range_name, safe="")
    )
    resp = session.get(url, timeout=60)
    if resp.status_code != 200:
        raise SystemExit(_brief(resp, spreadsheet_id, "Không đọc lại cột mã"))
    values = resp.json().get("values") or []
    column = []
    for row_index in range(n_rows):
        if row_index < len(values) and values[row_index]:
            column.append(values[row_index][0])
        else:
            column.append("")
    return column


def refresh_bridge(session, spreadsheet_id: str, rows: list[list]) -> dict:
    expect = ma_list(rows)
    if not expect:
        raise SystemExit("Nguồn không có SKU, không ghi cầu nối")
    from sync_products_from_excel import detect_columns

    ma_index = detect_columns(rows[0]).get("ma")
    if ma_index is None:
        raise SystemExit("Không thấy cột mã trên nguồn")
    sheets = _sheets_get(session, spreadsheet_id)
    plan = bridge_sheet_plan(sheets)
    _sheets_batch(session, spreadsheet_id, _delete_requests(plan["delete_before"], plan["keep"]))
    width = max((len(row) for row in rows), default=1)
    if plan["reuse"] is None:
        added = _sheets_batch(
            session,
            spreadsheet_id,
            [
                {
                    "addSheet": {
                        "properties": {
                            "title": TEMP_TITLE,
                            "gridProperties": {
                                "rowCount": max(len(rows) + 5, 20),
                                "columnCount": max(width + 2, 10),
                            },
                        }
                    }
                }
            ],
        )
        props = (((added.get("replies") or [{}])[0].get("addSheet") or {}).get("properties") or {})
        new_id = props.get("sheetId")
        if new_id is None:
            raise SystemExit("Không tạo được tab tạm trên cầu nối")
    else:
        new_id = plan["reuse"]
    _write_values(session, spreadsheet_id, TEMP_TITLE, rows)
    got = ma_from_column(_read_ma_column(session, spreadsheet_id, TEMP_TITLE, ma_index, len(rows)))
    if got != expect:
        if plan["reuse"] is None:
            _sheets_batch(session, spreadsheet_id, _delete_requests([new_id], plan["keep"]))
        raise SystemExit(f"Cầu nối lệch số SKU: nguồn {len(expect)}, sheet {len(got)}")
    rename = {
        "updateSheetProperties": {
            "properties": {"sheetId": new_id, "title": DANH_MUC_TITLE},
            "fields": "title",
        }
    }
    _sheets_batch(
        session,
        spreadsheet_id,
        _delete_requests(plan["delete_after"], plan["keep"]) + [rename],
    )
    after = _sheets_get(session, spreadsheet_id)
    after_keep = [row["sheetId"] for row in after if row.get("title") == CHINH_SACH_TITLE]
    if set(after_keep) != set(plan["keep"]):
        raise SystemExit("Tab Chính Sách Đối Tác bị đổi")
    if not any(row.get("title") == DANH_MUC_TITLE and row.get("sheetId") == new_id for row in after):
        raise SystemExit("Chưa đổi tên tab Danh Mục Vật Tư")
    return {"sku": len(expect), "delta": 0, "lastMa": expect[-1]}


def run_sync(xlsx: Path, title: str) -> dict:
    proc = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "sync_products_from_excel.py"),
            str(xlsx),
            "--drive-title",
            title,
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
    )
    if proc.returncode != 0:
        detail = ((proc.stdout or "") + "\n" + (proc.stderr or "")).strip()
        raise SystemExit(detail[-4000:] or f"sync thoát {proc.returncode}")
    found = None
    for candidate in (proc.stdout or "").splitlines():
        if not candidate.startswith("{"):
            continue
        try:
            obj = json.loads(candidate)
        except json.JSONDecodeError:
            continue
        if isinstance(obj, dict) and "excludedGroups" in obj:
            found = obj
    if not found:
        raise SystemExit("sync không in excludedGroups")
    return found


def run_checks() -> None:
    for name in ("check_rules.py", "check_seo.py"):
        proc = subprocess.run([sys.executable, str(ROOT / "scripts" / name)], cwd=ROOT)
        if proc.returncode != 0:
            raise SystemExit(f"{name} không đạt")


def git_changed_paths(root: Path) -> list[str]:
    existing = [item for item in GIT_ADD if (root / item).exists()]
    if not existing:
        return []
    subprocess.run(["git", "reset", "-q"], cwd=root, check=True)
    subprocess.run(["git", "add", "-A", "--", *existing], cwd=root, check=True)
    out = subprocess.check_output(["git", "diff", "--cached", "--name-only", "-z"], cwd=root)
    paths = [item.decode() for item in out.split(b"\0") if item]
    refused = [path for path in paths if not path_allowed(path)]
    if refused:
        subprocess.run(["git", "reset", "-q"], cwd=root, check=True)
        raise SystemExit("Từ chối commit: " + ", ".join(refused))
    return paths


def git_commit(root: Path, message: str) -> str:
    subprocess.run(
        [
            "git",
            "-c",
            "user.name=github-actions[bot]",
            "-c",
            "user.email=41898282+github-actions[bot]@users.noreply.github.com",
            "commit",
            "-m",
            message,
        ],
        cwd=root,
        check=True,
    )
    return subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=root, text=True).strip()


def git_push(root: Path, token: str, repo: str) -> None:
    url = f"https://x-access-token:{token}@github.com/{repo}.git"
    proc = subprocess.run(["git", "push", url, "HEAD:main"], cwd=root, capture_output=True, text=True)
    if proc.returncode != 0:
        err = redact((proc.stderr or "") + (proc.stdout or ""), [token])
        raise SystemExit("git push lỗi\n" + err[-2000:])


def list_pages_runs(token: str, repo: str) -> list[dict]:
    env = os.environ.copy()
    env["GH_TOKEN"] = token
    env["GITHUB_TOKEN"] = token
    errors = []
    for workflow in ("pages-build-deployment", "dynamic/pages/pages-build-deployment"):
        proc = subprocess.run(
            [
                "gh",
                "run",
                "list",
                "--repo",
                repo,
                "--workflow",
                workflow,
                "--limit",
                "20",
                "--json",
                "databaseId,status,conclusion,headSha,event,url",
            ],
            capture_output=True,
            text=True,
            env=env,
        )
        if proc.returncode != 0:
            errors.append(redact((proc.stderr or "") + (proc.stdout or ""), [token])[-500:])
            continue
        data = json.loads(proc.stdout or "[]")
        if not isinstance(data, list):
            raise SystemExit("Danh sách lượt Pages không đọc được")
        return data
    raise SystemExit("Không đọc được lượt Pages\n" + "\n".join(errors)[-1000:])


def main() -> int:
    import time

    from sync_schedule import wait_for_pages

    sa_json = os.environ.get("GDRIVE_SA_JSON") or ""
    bridge_id = (os.environ.get("LIVE_BRIDGE_SHEET_ID") or "").strip()
    token = (os.environ.get("GH_PUSH_TOKEN") or "").strip()
    repo = (os.environ.get("GITHUB_REPOSITORY") or "").strip()
    work = Path(os.environ.get("RUNNER_TEMP") or tempfile.mkdtemp(prefix="vp-sync-"))
    holder: dict = {}

    def metadata() -> dict:
        info = parse_sa(sa_json)
        holder["session"] = holder.get("session") or _session(info)
        holder["meta"] = fetch_metadata(holder["session"])
        return holder["meta"]

    def download(dest: Path) -> None:
        download_source(holder["session"], holder["meta"], dest)

    def bridge(path: Path) -> dict:
        return refresh_bridge(holder["session"], bridge_id, read_matrix(path))

    def push() -> None:
        if not repo:
            raise SystemExit("Thiếu GITHUB_REPOSITORY")
        git_push(ROOT, token, repo)

    def wait_pages(sha: str) -> None:
        if not repo:
            raise SystemExit("Thiếu GITHUB_REPOSITORY")
        wait_for_pages(lambda: list_pages_runs(token, repo), sha, time.sleep, time.time)

    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    job = Job(
        root=ROOT,
        state_path=ROOT / "scripts" / "sync_state.json",
        workdir=work,
        sa_json=sa_json,
        bridge_id=bridge_id,
        push_token=token,
        repo=repo,
        now=now,
        metadata=metadata,
        download=download,
        sync=run_sync,
        check=run_checks,
        bridge=bridge,
        changed_paths=lambda: git_changed_paths(ROOT),
        commit=lambda paths, message: git_commit(ROOT, message),
        push=push,
        wait_pages=wait_pages,
        log=print,
    )
    return run_job(job)


if __name__ == "__main__":
    sys.exit(main())
